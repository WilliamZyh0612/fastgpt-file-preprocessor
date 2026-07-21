from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .config import Settings


def _value(field, unknown): return field.value if field else unknown
def _evidence(field): return "；".join(f"{item.source} {item.location}: {item.excerpt}" for item in (field.evidence if field else []))
def _save(path: Path, headers: list[str], rows: list[list[str]]) -> None:
    book = Workbook(); sheet = book.active; sheet.title = "数据"; sheet.append(headers)
    for row in rows: sheet.append(row)
    for cell in sheet[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78"); cell.alignment = Alignment(horizontal="center", wrap_text=True)
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions; sheet.sheet_view.showGridLines = False
    for column in sheet.columns:
        letter = column[0].column_letter; sheet.column_dimensions[letter].width = min(52, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
        for cell in column: cell.alignment = Alignment(vertical="top", wrap_text=True)
    book.save(path)


def write_reports(records, output: Path, settings: Settings) -> None:
    output.mkdir(parents=True, exist_ok=True); unknown = settings.unknown_value
    catalog = [[r.filename, r.path, r.extension, r.sha256, r.normalized_text_hash, "、".join(c.name for c in r.categories) or unknown, "、".join(m.value for m in r.machine_models) or unknown, "、".join(x["brand"] for x in r.cnc_systems) or unknown, _value(r.version, unknown), _value(r.release_date, unknown), _value(r.visibility, unknown), _value(r.summary, unknown), "；".join(x.value for x in r.knowledge_points), "、".join(r.keywords), r.archive_suggestion, r.review_status, r.extraction_status, "；".join(r.flags), "；".join(r.errors)] for r in records]
    _save(output / "FastGPT知识库总目录.xlsx", ["文件名", "原始路径", "类型", "SHA-256", "标准化文本哈希", "文件类别", "适用机型", "数控系统品牌", "版本", "日期", "可见范围", "摘要", "核心知识点", "关键词", "建议归档目录", "审核状态", "提取状态", "风险标记", "错误原因"], catalog)
    def trail(record, item):
        refs = {e.evidence_id:e for e in record.evidence}; cited = [refs[x] for x in item.get("evidence_refs",[]) if x in refs]
        return [record.filename, "；".join(f"{e.location}" for e in cited), "；".join(e.excerpt for e in cited), item.get("confidence",0), record.review_status, "；".join(str(x) for x in record.conflicts) or "无"]
    products=[]; faults=[]; parts=[]; rules=[]
    for r in records:
        for x in r.ai_structured.get("product_records",[]): products.append([x[k] for k in ["product_model","product_name","processing_object","processing_range","key_parameters","accuracy","standard_configuration","optional_configuration","applicable_scenarios","limitations","cnc_system"]]+trail(r,x))
        for x in r.ai_structured.get("fault_records",[]): faults.append([x[k] for k in ["machine_model","cnc_system","fault_symptom","alarm_code","operating_condition","possible_causes","troubleshooting_steps","solution","safety_warning","requires_engineer"]]+trail(r,x))
        for x in r.ai_structured.get("part_records",[]): parts.append([x[k] for k in ["part_name","part_model","drawing_number","applicable_machine_models","installation_position","technical_specification","replacement_parts","replacement_limits"]]+trail(r,x))
        for x in r.ai_structured.get("business_rule_records",[]): rules.append([x[k] for k in ["business_module","rule_name","applicable_roles","trigger_conditions","allowed_actions","prohibited_actions","approval_requirements","status_transition","exception_handling"]]+trail(r,x))
    tail=["原始文件","页码/工作表/段落","原文证据","AI 置信度","审核状态","冲突标记"]
    _save(output / "产品选型矩阵.xlsx", ["产品型号","产品名称","加工对象","加工范围","关键参数","精度","标准配置","可选配置","适用场景","限制条件","数控系统",*tail], products)
    _save(output / "故障案例库.xlsx", ["机床型号","数控系统","故障现象","报警号","发生工况","可能原因","排查步骤","处理方法","安全提示","是否需人工售后",*tail], faults)
    _save(output / "配件适配表.xlsx", ["配件名称","配件型号","图号","适用机型","安装部位","技术规格","替代关系","替代限制",*tail], parts)
    _save(output / "CRM_ERP业务规则.xlsx", ["业务模块","规则名称","适用角色","触发条件","允许操作","禁止操作","审批要求","状态流转","异常处理",*tail], rules)
