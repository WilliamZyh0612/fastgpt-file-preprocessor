from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

UNKNOWN = "待人工确认"
SUPPORTED = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".xlsm", ".ppt", ".pptx", ".txt", ".md", ".markdown"}
STOPWORDS = {"的", "和", "及", "与", "在", "是", "为", "了", "本", "该", "产品", "文件", "说明", "manual", "the", "and", "for", "with", "this"}


@dataclass
class Record:
    path: str
    filename: str
    extension: str
    sha256: str
    size_bytes: int
    category: str = UNKNOWN
    machine_model: str = UNKNOWN
    cnc_system: str = UNKNOWN
    version: str = UNKNOWN
    release_date: str = UNKNOWN
    visibility: str = UNKNOWN
    summary: str = UNKNOWN
    knowledge_points: str = UNKNOWN
    keywords: str = UNKNOWN
    archive_suggestion: str = UNKNOWN
    extraction_status: str = "成功"
    flags: list[str] = field(default_factory=list)
    text: str = field(default="", repr=False)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix == ".docx":
        from docx import Document
        doc = Document(path)
        blocks = [p.text for p in doc.paragraphs]
        blocks.extend(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
        return "\n".join(blocks)
    if suffix == ".doc":
        return _extract_legacy_office(path, "word")
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        blocks = []
        for sheet in workbook.worksheets:
            blocks.append(f"工作表：{sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                blocks.append(" | ".join(str(v) for v in row if v is not None))
        return "\n".join(blocks)
    if suffix == ".xls":
        return _extract_legacy_office(path, "excel")
    if suffix == ".pptx":
        from pptx import Presentation
        presentation = Presentation(path)
        return "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text"))
    if suffix == ".ppt":
        return _extract_legacy_office(path, "powerpoint")
    raise ValueError(f"不支持的文件类型：{suffix}")


def _extract_legacy_office(path: Path, application: str) -> str:
    """Read legacy Office files through a locally installed Office application, never saving them."""
    try:
        import win32com.client  # pywin32 + locally installed Microsoft Office are required
    except ImportError as exc:
        raise RuntimeError("旧版 Office 文件需要安装 Microsoft Office 和 pywin32") from exc
    if application == "word":
        app = win32com.client.DispatchEx("Word.Application"); app.Visible = False
        try:
            document = app.Documents.Open(str(path.resolve()), ReadOnly=True, AddToRecentFiles=False)
            try: return document.Content.Text
            finally: document.Close(False)
        finally: app.Quit()
    if application == "excel":
        app = win32com.client.DispatchEx("Excel.Application"); app.Visible = False
        try:
            book = app.Workbooks.Open(str(path.resolve()), ReadOnly=True, UpdateLinks=0)
            try:
                return "\n".join(" | ".join(str(value) for value in row if value is not None) for sheet in book.Worksheets for row in (sheet.UsedRange.Value if isinstance(sheet.UsedRange.Value, tuple) else ((sheet.UsedRange.Value,),)))
            finally: book.Close(False)
        finally: app.Quit()
    app = win32com.client.DispatchEx("PowerPoint.Application")
    try:
        presentation = app.Presentations.Open(str(path.resolve()), ReadOnly=True, Untitled=False, WithWindow=False)
        try:
            return "\n".join(shape.TextFrame.TextRange.Text for slide in presentation.Slides for shape in slide.Shapes if shape.HasTextFrame and shape.TextFrame.HasText)
        finally: presentation.Close()
    finally: app.Quit()


def first_match(patterns: Iterable[str], text: str) -> str:
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return match.group(1).strip()
    return UNKNOWN


def infer_category(text: str) -> str:
    lower = text.lower()
    rules = [("故障案例", ("故障", "报警", "维修", "异常", "error", "alarm")), ("配件资料", ("配件", "备件", "零件", "spare part")),
             ("产品资料", ("选型", "参数", "规格", "product", "型号")), ("业务规则", ("crm", "erp", "报价", "订单", "客户", "审批"))]
    return next((name for name, words in rules if any(word in lower for word in words)), UNKNOWN)


def keywords(text: str) -> str:
    words = re.findall(r"[A-Za-z][A-Za-z0-9+._-]{2,}|[\u4e00-\u9fff]{2,}", text)
    counts = Counter(word for word in words if word.lower() not in STOPWORDS)
    return "、".join(word for word, _ in counts.most_common(8)) or UNKNOWN


def short_sentences(text: str) -> list[str]:
    clean = re.sub(r"\s+", " ", text).strip()
    return [part.strip() for part in re.split(r"[。！？.!?]\s*", clean) if len(part.strip()) >= 12]


def analyze(path: Path) -> Record:
    record = Record(path=str(path.resolve()), filename=path.name, extension=path.suffix.lower(), sha256=_sha256(path), size_bytes=path.stat().st_size)
    try:
        record.text = extract_text(path)
        if not record.text.strip():
            record.extraction_status, record.flags = "需人工处理", ["未提取到可分析文本"]
            return record
        sentences = short_sentences(record.text)
        record.category = infer_category(record.text)
        record.machine_model = first_match((r"(?:机床型号|型号|model)\s*[:：#]?\s*([A-Z]{1,5}[- ]?\d{2,5}[A-Z0-9-]*)",), record.text)
        record.cnc_system = first_match((r"(?:数控系统|CNC系统|控制系统)\s*[:：]?\s*([A-Za-z]+\s*[0-9A-Za-z-]+)", r"\b(FANUC\s*[0-9A-Za-z-]*|SIEMENS\s*[0-9A-Za-z-]*|Mitsubishi\s*[0-9A-Za-z-]*)\b"), record.text)
        record.version = first_match((r"(?:版本|version|ver\.)\s*[:：]?\s*(v?\d+(?:\.\d+){0,3}(?:[-_A-Za-z0-9]+)?)",), record.text)
        record.release_date = first_match((r"(?:发布日期|发布|日期|date)\s*[:：]?\s*((?:20\d{2}|19\d{2})[-/.年]\d{1,2}[-/.月]\d{1,2}日?)",), record.text)
        record.visibility = first_match((r"(?:可见范围|保密级别|visibility)\s*[:：]?\s*(内部|公开|销售|售后|管理层|客户|public|internal)",), record.text)
        record.summary = (sentences[0][:240] if sentences else UNKNOWN)
        record.knowledge_points = "；".join(sentences[:5])[:1000] if sentences else UNKNOWN
        record.keywords = keywords(record.text)
        category = record.category if record.category != UNKNOWN else "待分类"
        model = record.machine_model if record.machine_model != UNKNOWN else "待确认型号"
        record.archive_suggestion = f"/{category}/{model}/{record.release_date if record.release_date != UNKNOWN else '待确认日期'}"
    except Exception as exc:  # preserve source and expose the failure for audit
        record.extraction_status, record.flags = "需人工处理", [f"提取失败：{type(exc).__name__}"]
    return record


def _tokens(text: str) -> set[str]:
    return {x.lower() for x in re.findall(r"[A-Za-z0-9\u4e00-\u9fff]{2,}", text) if x.lower() not in STOPWORDS}


def enrich(records: list[Record]) -> None:
    hashes: dict[str, list[Record]] = {}
    for item in records:
        hashes.setdefault(item.sha256, []).append(item)
    for group in hashes.values():
        if len(group) > 1:
            for item in group: item.flags.append("重复文件")
    for index, current in enumerate(records):
        current_tokens = _tokens(current.text)
        for other in records[index + 1:]:
            other_tokens = _tokens(other.text)
            score = len(current_tokens & other_tokens) / max(1, len(current_tokens | other_tokens))
            if score >= 0.72 and current.sha256 != other.sha256:
                current.flags.append(f"相似文件：{other.filename}（{score:.0%}）")
                other.flags.append(f"相似文件：{current.filename}（{score:.0%}）")
    groups: dict[tuple[str, str], list[Record]] = {}
    for item in records:
        if item.machine_model != UNKNOWN:
            groups.setdefault((item.category, item.machine_model), []).append(item)
    for group in groups.values():
        versions = {item.version for item in group if item.version != UNKNOWN}
        if len(versions) > 1:
            for item in group: item.flags.append("版本冲突：同类同型号存在多个版本")
        dates = sorted(item.release_date for item in group if item.release_date != UNKNOWN)
        if len(dates) > 1:
            for item in group:
                if item.release_date == dates[0]: item.flags.append("疑似过期资料：存在较新发布日期")
    for item in records:
        if not item.flags: item.flags.append("无")


def write_reports(records: list[Record], destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); wb.remove(wb.active)
    columns = ["文件名", "原始路径", "类型", "SHA-256", "分类", "机床型号", "数控系统", "版本", "发布日期", "可见范围", "摘要", "核心知识点", "关键词", "建议归档目录", "处理状态", "风险标记"]
    rows = [[r.filename, r.path, r.extension, r.sha256, r.category, r.machine_model, r.cnc_system, r.version, r.release_date, r.visibility, r.summary, r.knowledge_points, r.keywords, r.archive_suggestion, r.extraction_status, "；".join(r.flags)] for r in records]
    sheets = {
        "知识库总目录": (columns, rows),
        "产品选型矩阵": (["机床型号", "数控系统", "版本", "资料", "摘要", "待确认项"], [[r.machine_model, r.cnc_system, r.version, r.filename, r.summary, "是" if UNKNOWN in (r.machine_model, r.cnc_system, r.version) else "否"] for r in records if r.category == "产品资料"]),
        "故障案例库": (["机床型号", "数控系统", "文件", "核心知识点", "风险标记"], [[r.machine_model, r.cnc_system, r.filename, r.knowledge_points, "；".join(r.flags)] for r in records if r.category == "故障案例"]),
        "配件适配表": (["机床型号", "数控系统", "文件", "关键词", "待确认项"], [[r.machine_model, r.cnc_system, r.filename, r.keywords, "是" if UNKNOWN in (r.machine_model, r.cnc_system) else "否"] for r in records if r.category == "配件资料"]),
        "CRM_ERP业务规则": (["文件", "摘要", "核心知识点", "可见范围", "待确认项"], [[r.filename, r.summary, r.knowledge_points, r.visibility, "是" if r.visibility == UNKNOWN else "否"] for r in records if r.category == "业务规则"]),
    }
    for name, (headers, data) in sheets.items():
        sheet = wb.create_sheet(name); sheet.append(headers)
        for row in data: sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78"); cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = min(48, max(14, max(len(str(c.value or "")) for c in col) + 2))
            for cell in col: cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(destination)


def run(source: Path, output: Path) -> list[Record]:
    files = [p for p in source.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED and not p.resolve().is_relative_to(output.resolve())]
    records = [analyze(path) for path in files]
    enrich(records); output.mkdir(parents=True, exist_ok=True)
    write_reports(records, output / "FastGPT知识库总目录.xlsx")
    (output / "analysis.json").write_text(json.dumps([asdict(r) | {"text": None} for r in records], ensure_ascii=False, indent=2), encoding="utf-8")
    return records


def main() -> None:
    parser = argparse.ArgumentParser(description="FastGPT 知识库文件预处理器（只读源文件）")
    parser.add_argument("--input", required=True, type=Path, help="待处理资料目录")
    parser.add_argument("--output", required=True, type=Path, help="分析结果目录，必须不在输入目录内")
    args = parser.parse_args()
    if not args.input.is_dir(): parser.error("--input 必须是存在的目录")
    if args.output.resolve().is_relative_to(args.input.resolve()): parser.error("--output 不得位于 --input 内，避免将生成物再次扫描")
    records = run(args.input, args.output)
    print(f"完成：处理 {len(records)} 个文件；输出：{args.output.resolve()}")


if __name__ == "__main__": main()
