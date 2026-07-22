from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path

from .config import Settings

SUPPORTED = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".xlsm", ".ppt", ".pptx", ".txt", ".md", ".markdown"}


@dataclass
class Evidence:
    evidence_id: str
    source: str
    location: str
    excerpt: str


@dataclass
class Extracted:
    text: str
    evidence: list[Evidence] = field(default_factory=list)
    tables: list[dict] = field(default_factory=list)
    headings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    ocr_attempted: bool = False
    ocr_status: str = "未触发"
    ocr_model: str = ""
    chunks: list[list[Evidence]] = field(default_factory=list)


def evidence(source: str, location: str, excerpt: str) -> Evidence:
    token = hashlib.sha256(f"{source}|{location}|{excerpt}".encode("utf-8")).hexdigest()[:16]
    return Evidence(token, source, location, excerpt)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for part in iter(lambda: source.read(1024 * 1024), b""): digest.update(part)
    return digest.hexdigest()


def extract(path: Path, settings: Settings) -> Extracted:
    if path.stat().st_size > settings.max_file_size_mb * 1024 * 1024:
        return Extracted("", errors=[f"文件超过限制：{settings.max_file_size_mb}MB"])
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        items=[evidence(path.name, f"正文分块 {i//settings.chunk_size_chars+1}", text[i:i+settings.chunk_size_chars]) for i in range(0,len(text),max(1,settings.chunk_size_chars-settings.chunk_overlap_chars))] or [evidence(path.name,"正文","")]
        result = Extracted(text, items); result.chunks = [[item] for item in items]; return result
    if suffix == ".pdf": return _pdf(path, settings)
    if suffix == ".docx": return _docx(path, settings)
    if suffix in {".xlsx", ".xlsm"}: return _xlsx(path, settings)
    if suffix == ".pptx": return _pptx(path, settings)
    if suffix in {".doc", ".xls", ".ppt"}:
        if not settings.legacy_office_conversion_enabled: return Extracted("", errors=["旧版 Office 文件需安全转换"], ocr_attempted=False)
        from .legacy_office import extract_legacy_readonly
        text = extract_legacy_readonly(path, settings.legacy_office_timeout_seconds)
        items = [evidence(path.name, f"安全转换正文 {index // max(1, settings.chunk_size_chars - settings.chunk_overlap_chars) + 1}", text[index:index + settings.chunk_size_chars]) for index in range(0, len(text), max(1, settings.chunk_size_chars - settings.chunk_overlap_chars))] or [evidence(path.name, "安全转换正文", "")]
        result = Extracted(text, items); result.chunks = [[item] for item in items]; return result
    return Extracted("", errors=[f"不支持的文件类型：{suffix}"])


def _pdf(path: Path, settings: Settings) -> Extracted:
    from pypdf import PdfReader
    result = Extracted("")
    reader = PdfReader(path)
    pages = []
    for index, page in enumerate(reader.pages, 1):
        content = page.extract_text() or ""; pages.append(content)
        if content: result.evidence.append(evidence(path.name, f"PDF 第{index}页", content))
    result.text = "\n".join(pages)
    if len(result.text.strip()) < settings.min_pdf_text_chars:
        result.ocr_attempted = True; result.ocr_status = "待视觉识别"
        result.errors.append("PDF 有效文字过少；视觉 OCR 未配置或未完成，未提取 OCR 文字")
    result.chunks = _chunk_evidence(result.evidence, settings.chunk_size_chars, settings.chunk_overlap_chars); return result


def _docx(path: Path, settings: Settings) -> Extracted:
    from docx import Document
    document = Document(path); result = Extracted(""); pieces = []
    for i, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text.strip()
        if not text: continue
        location = f"段落 {i}"
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            location = f"标题层级 {paragraph.style.name}"; result.headings.append(text)
        pieces.append(text); result.evidence.append(evidence(path.name, location, text))
    for index, table in enumerate(document.tables, 1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        result.tables.append({"source": path.name, "location": f"Word 表格 {index}", "rows": rows})
        pieces.extend(" | ".join(row) for row in rows)
    result.text = "\n".join(pieces); result.chunks = _chunk_evidence(result.evidence, settings.chunk_size_chars, settings.chunk_overlap_chars); return result


def _xlsx(path: Path, settings: Settings) -> Extracted:
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True); result = Extracted(""); pieces = []
    try:
        for sheet in book.worksheets:
            rows = [["" if value is None else str(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            result.tables.append({"source": path.name, "location": f"工作表 {sheet.title}", "rows": rows})
            for number, row in enumerate(rows, 1):
                line = " | ".join(row).strip()
                if line: pieces.append(line); result.evidence.append(evidence(path.name, f"工作表 {sheet.title} 第{number}行", line))
        result.text = "\n".join(pieces); result.chunks = _chunk_evidence(result.evidence, settings.chunk_size_chars, settings.chunk_overlap_chars); return result
    finally:
        book.close()


def _pptx(path: Path, settings: Settings) -> Extracted:
    from pptx import Presentation
    result = Extracted(""); pieces = []
    for number, slide in enumerate(Presentation(path).slides, 1):
        values = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        text = "\n".join(values)
        if text: pieces.append(text); result.evidence.append(evidence(path.name, f"PPT 第{number}页", text))
    result.text = "\n".join(pieces); result.chunks = _chunk_evidence(result.evidence, settings.chunk_size_chars, settings.chunk_overlap_chars); return result


def _chunk_evidence(items: list[Evidence], size: int, overlap: int = 0) -> list[list[Evidence]]:
    """Split by actual character count, retaining an exact textual overlap.

    Evidence may originate from a page, heading, row or table cell.  When a
    boundary cuts one of these sources, a new evidence fragment retains its
    source and location.  Therefore every AI request remains auditable while
    no request exceeds ``size`` characters.
    """
    if size < 1 or overlap < 0 or overlap >= size: raise ValueError("分块大小或重叠长度不合法")

    def tail(group: list[Evidence]) -> list[Evidence]:
        remaining = overlap; selected: list[Evidence] = []
        for item in reversed(group):
            if remaining <= 0: break
            text = item.excerpt[-remaining:]
            selected.append(evidence(item.source, f"{item.location} 重叠片段", text))
            remaining -= len(text)
        return list(reversed(selected))

    groups: list[list[Evidence]] = []; current: list[Evidence] = []; used = 0
    for original in items:
        text = original.excerpt; start = 0; part = 1
        while start < len(text) or (not text and not current and not groups):
            if used == size:
                groups.append(current); current = tail(current); used = sum(len(item.excerpt) for item in current)
            take = min(size - used, len(text) - start)
            if take <= 0: break
            excerpt = text[start:start + take]
            location = original.location if start == 0 and take == len(text) else f"{original.location} 片段 {part}"
            current.append(original if location == original.location else evidence(original.source, location, excerpt))
            used += take; start += take; part += 1
            if used == size and start < len(text):
                groups.append(current); current = tail(current); used = sum(len(item.excerpt) for item in current)
    if current: groups.append(current)
    return groups
