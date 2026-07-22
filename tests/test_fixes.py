import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from dataclasses import asdict
from pathlib import Path

from openpyxl import load_workbook

from src.ai import AI_SCHEMA, AIAnalyzer, validate_ai_result
from src.config import Settings, load_settings
from src.parsers import _chunk_evidence, evidence, extract
from src.preprocessor import Record, analyze, enrich
from src.reports import write_reports


def structured_result(evidence_id: str, product_name: str = "BK5030 选型") -> dict:
    def field(value: str) -> dict:
        return {"value": value, "confidence": .9, "evidence_refs": [evidence_id]}
    metadata = {
        "categories": [{"name": "产品资料", "confidence": .9, "reason": "参数表", "evidence_refs": [evidence_id]}],
        "machine_models": [field("BK5030")],
        "cnc_systems": [{"brand": "FANUC", "model": "0i-MF", "version": "待人工确认", "confidence": .8, "evidence_refs": [evidence_id]}],
        "visibility": field("内部"), "version": field("V1.2"), "release_date": field("2025-01-01"),
        "summary": field("BK5030 产品参数说明"), "knowledge_points": [field("加工范围 500mm")], "keywords": [field("BK5030")],
    }
    product = {"product_model": "BK5030", "product_name": product_name, "processing_object": "齿轮", "processing_range": "500mm", "key_parameters": "行程 500mm", "accuracy": "0.01mm", "standard_configuration": "标准夹具", "optional_configuration": "自动上下料", "applicable_scenarios": "批量加工", "limitations": "需人工确认", "cnc_system": "FANUC 0i-MF", "evidence_refs": [evidence_id], "confidence": .9}
    return {"document_metadata": metadata, "product_records": [product], "fault_records": [], "part_records": [], "business_rule_records": []}


class FixTests(unittest.TestCase):
    def test_strict_schema_is_complete_and_payload_is_valid(self):
        schema = AI_SCHEMA["schema"]
        self.assertEqual(set(schema["required"]), set(schema["properties"]))
        def inspect(node):
            if not isinstance(node, dict): return
            if node.get("type") == "object":
                self.assertTrue(node.get("properties"))
                self.assertEqual(set(node["required"]), set(node["properties"]))
                self.assertIs(node["additionalProperties"], False)
            if node.get("type") == "array": self.assertIn("items", node)
            for value in node.values():
                if isinstance(value, dict): inspect(value)
                elif isinstance(value, list):
                    for item in value: inspect(item)
        inspect(schema)
        payload = AIAnalyzer(Settings(ai_enabled=True, api_base_url="http://x/v1", text_model="m"))._payload("x", [])
        self.assertEqual(payload["response_format"], {"type": "json_schema", "json_schema": AI_SCHEMA})

    def test_last_chunk_and_configured_overlap_are_effective(self):
        parts = [evidence("x", str(i), "a" * 8) for i in range(3)]
        chunks = _chunk_evidence(parts, 12, 4)
        self.assertEqual(chunks[-1][-1].evidence_id, parts[-1].evidence_id)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "long.txt"; path.write_text("a" * 1200, encoding="utf-8")
            config = Path(tmp) / "config.json"; config.write_text(json.dumps({"chunk_size_chars": 500, "chunk_overlap_chars": 100}), encoding="utf-8")
            parsed = extract(path, load_settings(config))
            self.assertGreater(len(parsed.chunks), 2)
            self.assertEqual(parsed.chunks[0][0].excerpt[-100:], parsed.chunks[1][0].excerpt[:100])

    def test_long_txt_uses_isolated_chunk_evidence_and_aggregates_records(self):
        calls: list[dict] = []
        def transport(payload):
            source = json.loads(payload["messages"][1]["content"]); calls.append(source)
            return {"choices": [{"message": {"content": json.dumps(structured_result(source["evidence"][0]["evidence_id"], f"记录-{len(calls)}"), ensure_ascii=False)}}]}
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "long.txt"; path.write_text("BK5030 产品参数。" * 200, encoding="utf-8")
            record = analyze(path, Settings(ai_enabled=True, api_base_url="http://x/v1", text_model="m", chunk_size_chars=500, chunk_overlap_chars=100), AIAnalyzer(Settings(ai_enabled=True, api_base_url="http://x/v1", text_model="m", chunk_size_chars=500, chunk_overlap_chars=100), transport))
        self.assertGreater(len(calls), 1)
        self.assertEqual(len(record.ai_structured["product_records"]), len(calls))
        for source in calls:
            self.assertEqual(len(source["evidence"]), 1)
            self.assertIn(source["evidence"][0]["excerpt"], source["context"])

    def test_ai_metadata_requires_current_nonempty_evidence(self):
        item = evidence("x", "第1页", "BK5030 产品参数")
        result = structured_result(item.evidence_id)
        validate_ai_result(result, {item.evidence_id})
        meta = result["document_metadata"]
        for field in ("visibility", "version", "release_date", "summary"):
            self.assertTrue(meta[field]["evidence_refs"])
        for field in ("machine_models", "knowledge_points", "keywords", "cnc_systems"):
            self.assertTrue(meta[field][0]["evidence_refs"])
        bad = structured_result(item.evidence_id); bad["document_metadata"]["summary"]["evidence_refs"] = []
        with self.assertRaises(ValueError): validate_ai_result(bad, {item.evidence_id})
        wrong = structured_result(item.evidence_id); wrong["document_metadata"]["version"]["evidence_refs"] = ["0" * 16]
        with self.assertRaises(ValueError): validate_ai_result(wrong, {item.evidence_id})

    def test_empty_text_is_not_standardized_text_duplicate(self):
        left = Record("a", "a.txt", ".txt", "a", 0); right = Record("b", "b.txt", ".txt", "b", 0)
        enrich([left, right], Settings())
        self.assertNotIn("标准化文本重复", left.flags)
        self.assertNotIn("标准化文本重复", right.flags)

    def test_professional_excel_uses_structured_record_fields(self):
        item = evidence("manual.pdf", "PDF 第2页", "BK5030 加工范围 500mm")
        record = Record("manual.pdf", "manual.pdf", ".pdf", "hash", 10)
        record.evidence = [item]; record.ai_structured = structured_result(item.evidence_id)["product_records"] and {"product_records": structured_result(item.evidence_id)["product_records"]}
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp); write_reports([record], output, Settings())
            book = load_workbook(output / "产品选型矩阵.xlsx", read_only=True); sheet = book.active
            row = list(sheet.iter_rows(min_row=2, max_row=2, values_only=True))[0]; book.close()
            self.assertEqual(row[0], "BK5030"); self.assertEqual(row[1], "BK5030 选型")
            self.assertEqual(row[3], "500mm"); self.assertIn("PDF 第2页", row[-5])

    def test_real_package_can_be_extracted_and_started(self):
        root = Path(__file__).parents[1]
        subprocess.run([sys.executable, str(root / "tools" / "package.py")], cwd=root, check=True)
        archive = root / "dist" / "fastgpt-file-preprocessor.zip"
        with tempfile.TemporaryDirectory() as tmp:
            with zipfile.ZipFile(archive) as bundle: bundle.extractall(tmp)
            package = Path(tmp) / "fastgpt-file-preprocessor"; input_dir = Path(tmp) / "input"; output = Path(tmp) / "output"; input_dir.mkdir()
            input_dir.joinpath("smoke.txt").write_text("BK5030 产品参数", encoding="utf-8")
            environment = {**os.environ, "PYTHONPATH": str(package)}
            subprocess.run([sys.executable, "-m", "src.preprocessor", "--input", str(input_dir), "--output", str(output), "--config", str(package / "config.json")], cwd=package, env=environment, check=True)
            self.assertTrue((output / "analysis.json").exists())


if __name__ == "__main__": unittest.main()
