import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from dataclasses import asdict

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas

from src.ai import AIAnalyzer, AIResponseError
from src.config import Settings, load_settings
from src.parsers import evidence, extract
from src.preprocessor import REVIEW_PENDING, analyze, run, split_sentences


class V020Tests(unittest.TestCase):
    def setUp(self): self.settings = Settings(similarity_threshold=.35, min_pdf_text_chars=30)

    def _root(self):
        temp = tempfile.TemporaryDirectory(); self.addCleanup(temp.cleanup)
        source = Path(temp.name) / "source"; source.mkdir(); return source, Path(temp.name) / "output"

    def test_supported_pdf_docx_xlsx_pptx_and_outputs(self):
        source, output = self._root()
        pdf = source / "BK5030_产品资料.pdf"; page = canvas.Canvas(str(pdf)); page.drawString(72, 720, "BK5030 product specification and operation manual"); page.save()
        document = Document(); document.add_heading("TCK56 操作说明", level=1); document.add_paragraph("广数 GSK 980 操作与维修资料。"); document.save(source / "TCK56.docx")
        book = Workbook(); sheet = book.active; sheet.title = "参数"; sheet.append(["型号", "系统"]); sheet.append(["DC260U", "FANUC 0i-MF"]); book.save(source / "DC260U.xlsx")
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text = "BS712-N 配件"; slide.placeholders[1].text = "备件型号 P-01"; deck.save(source / "BS712-N.pptx")
        records = run(source, output, self.settings)
        self.assertEqual(len(records), 4)
        self.assertTrue(any("DC-260U" in [m.value for m in r.machine_models] for r in records))
        for name in ["FastGPT知识库总目录.xlsx", "产品选型矩阵.xlsx", "故障案例库.xlsx", "配件适配表.xlsx", "CRM_ERP业务规则.xlsx", "analysis.json"]: self.assertTrue((output / name).exists())

    def test_scanned_pdf_falls_back_without_fabrication(self):
        source, _ = self._root(); path = source / "scan.pdf"; page = canvas.Canvas(str(path)); page.rect(20, 20, 200, 100); page.save()
        result = extract(path, self.settings)
        self.assertTrue(result.ocr_attempted); self.assertIn("未配置或未完成", result.errors[0])

    def test_chinese_models_cnc_and_filename_detection(self):
        source, _ = self._root(); path = source / "新代_DC260U_维修.txt"; path.write_text("广数系统故障处理", encoding="utf-8")
        record = analyze(path, self.settings)
        self.assertEqual([m.value for m in record.machine_models], ["DC-260U"])
        self.assertEqual({x["brand"] for x in record.cnc_systems}, {"广数", "新代"})

    def test_model_aliases_do_not_match_parent_models(self):
        source, _ = self._root()
        for text, expected in [("BS712-N 配件", "BS712-N"), ("BS712-R 配件", "BS712-R"), ("BS1018-B 参数", "BS1018-B"), ("DC 260U", "DC-260U"), ("Y5125B CNC", "Y5125B CNC")]:
            path=source / f"{expected}.txt"; path.write_text(text,encoding="utf-8")
            self.assertEqual([x.value for x in analyze(path,self.settings).machine_models],[expected])

    def test_multilabel_and_decimal_versions_are_not_sentence_split(self):
        source, _ = self._root(); path = source / "manual.txt"; path.write_text("BK5030 产品参数、操作说明、维修报警和配件清单。版本：V1.2。精度 0.01mm，圆度 3.14。", encoding="utf-8")
        record = analyze(path, self.settings); names = {item.name for item in record.categories}
        self.assertTrue({"产品资料", "操作资料", "维修资料", "配件资料"}.issubset(names)); self.assertEqual(record.version.value, "V1.2")
        self.assertIn("精度 0.01mm，圆度 3.14", split_sentences(path.read_text(encoding="utf-8")))

    def test_duplicates_similarity_version_conflict_and_source_unchanged(self):
        source, output = self._root(); first = source / "BK5030 V1.txt"; content = "BK5030 产品参数 版本：V1.0 发布日期：2024-01-01 加工范围 500mm"
        first.write_text(content, encoding="utf-8"); (source / "copy.txt").write_text(content, encoding="utf-8"); (source / "BK5030 V2.txt").write_text("BK5030 产品参数 版本：V2.0 发布日期：2025-01-01 加工范围 510mm", encoding="utf-8")
        records = run(source, output, self.settings)
        self.assertEqual(first.read_text(encoding="utf-8"), content)
        self.assertTrue(any("完全重复文件" in flag for record in records for flag in record.flags)); self.assertTrue(any("版本冲突" in flag for record in records for flag in record.flags))
        self.assertTrue(all(record.review_status == REVIEW_PENDING for record in records))

    def test_empty_and_damaged_files_and_output_not_scanned(self):
        source, output = self._root(); (source / "empty.txt").write_text("", encoding="utf-8"); (source / "broken.xlsx").write_bytes(b"not an xlsx")
        records = run(source, output, self.settings); self.assertEqual(len(records), 2)
        self.assertTrue(all(record.extraction_status == "需人工处理" for record in records))
        again = run(source, output, self.settings); self.assertEqual(len(again), 2)

    def test_config_applies_and_key_is_not_serialized(self):
        source, output = self._root(); (source / "a.txt").write_text("BK5030", encoding="utf-8")
        config = source.parent / "config.json"; config.write_text(json.dumps({"unknown_value": "需确认", "similarity_threshold": .5, "max_file_size_mb": 1}, ensure_ascii=False), encoding="utf-8")
        settings = load_settings(config); self.assertEqual(settings.unknown_value, "需确认")
        records = run(source, output, settings); self.assertEqual(records[0].visibility.value, "需确认")
        self.assertNotIn("FASTGPT_PREPROCESSOR_API_KEY", (output / "analysis.json").read_text(encoding="utf-8"))

    def test_ai_invalid_json_and_timeout_retry(self):
        calls = {"count": 0}
        def invalid(_): return {"choices": [{"message": {"content": "not-json"}}]}
        analyzer = AIAnalyzer(Settings(ai_enabled=True, api_base_url="http://test/v1", text_model="test", retries=1), invalid)
        with self.assertRaises(AIResponseError): analyzer.analyze("text")
        def delayed(_):
            calls["count"] += 1
            if calls["count"] < 2: raise TimeoutError("timeout")
            ref = json.loads(_["messages"][1]["content"])["evidence"][0]["evidence_id"]
            field = lambda value: {"value":value,"confidence":.8,"evidence_refs":[ref]}
            data={"document_metadata":{"categories":[],"machine_models":[],"cnc_systems":[],"visibility":field("待人工确认"),"version":field("待人工确认"),"release_date":field("待人工确认"),"summary":field("x"),"knowledge_points":[],"keywords":[]},"product_records":[],"fault_records":[],"part_records":[],"business_rule_records":[]}
            return {"choices": [{"message": {"content": json.dumps(data)}}]}
        source = evidence("test.txt", "正文", "text")
        self.assertEqual(AIAnalyzer(Settings(ai_enabled=True, api_base_url="http://test/v1", text_model="test", retries=1), delayed).analyze("text", [asdict(source)])["document_metadata"]["summary"]["value"], "x")
        self.assertEqual(calls["count"], 2)

    def test_legacy_office_requires_safe_conversion(self):
        source, _ = self._root(); file = source / "old.doc"; file.write_bytes(b"legacy")
        self.assertIn("需安全转换", extract(file, self.settings).errors[0])

    def test_legacy_safe_conversion_timeout_is_recorded(self):
        source, _ = self._root(); file = source / "slow.doc"; file.write_bytes(b"legacy")
        settings = Settings(legacy_office_conversion_enabled=True, legacy_office_timeout_seconds=1)
        with patch("src.legacy_office.extract_legacy_readonly", side_effect=TimeoutError("timeout")):
            record = analyze(file, settings)
        self.assertEqual(record.extraction_status, "需人工处理")
        self.assertIn("TimeoutError", record.errors[0])


if __name__ == "__main__": unittest.main()
